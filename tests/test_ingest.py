"""Čtení dokumentů: značky členění, skládání řádků, rozbor na bloky, Word, Markdown a složka.

Všechny texty jsou smyšlené (město Horní Áíčkov). Soubory Wordu si testy vyrábějí samy.
"""
from __future__ import annotations

import json

import pytest

from chatbot.ingest import markers
from chatbot.ingest.extract import NoTextError, read_document
from chatbot.ingest.folder import doc_id, load_documents
from chatbot.ingest.model import Line
from chatbot.ingest.structure import build_blocks, join_wrapped


def L(text: str, **kwargs) -> Line:
    return Line(text=text, **kwargs)


# ------------------------------------------------------------------ značky
@pytest.mark.parametrize("text, expected", [
    ("Čl. 3", ("3", "")),
    ("Článek 4.", ("4", "")),
    ("ČLÁNEK IV", ("4", "")),
    ("Čl. 2 Doba nočního klidu", ("2", "Doba nočního klidu")),
    ("Čl. 5a", ("5a", "")),
    ("§ 3", ("§ 3", "")),
    ("§ 12a", ("§ 12a", "")),
    ("§ 3 Sazba poplatku", ("§ 3", "Sazba poplatku")),      # starší vyhlášky a zákony se člení na paragrafy
])
def test_nadpis_clanku(text, expected):
    assert markers.article(text) == expected


@pytest.mark.parametrize("text", [
    "čl. 4 odst. 2 této vyhlášky.",         # odkaz uprostřed věty, který se zalomil na začátek řádku
    "Čl. 4 odst. 2 se použije obdobně",
    "§ 84 odst. 2 písm. h) zákona č. 128/2000 Sb., o obcích",
    "§ 14 zákona č. 565/1990 Sb., o místních poplatcích",
    "§ 4 se zrušuje",
    "§ 2 Zákona o místních poplatcích se použije takto:",
    "Článek o nočním klidu vyšel v novinách.",
])
def test_odkaz_na_clanek_neni_nadpis(text):
    assert markers.article(text) is None


PARAGRAFY = """Vyhláška města Horní Áíčkov o místních poplatcích

§ 1
Předmět úpravy

(1) Město Horní Áíčkov vybírá místní poplatek ze psů.
(2) Správcem poplatku je městský úřad.

§ 2
Poplatek platí držitel psa.

§ 3 Účinnost

Tato vyhláška nabývá účinnosti dnem 1. ledna 2026.
"""


def test_predpis_cleneny_na_paragrafy(tmp_path):
    """„§ 1“ s názvem na dalším řádku, bez názvu (hned text) i s názvem na témže řádku. Citace nese paragraf."""
    from chatbot.index.chunks import make_chunks

    (tmp_path / "vyhlaska.txt").write_text(PARAGRAFY, encoding="utf-8")
    docs, skipped = load_documents(tmp_path)
    assert not skipped and docs[0].flags == []
    assert [(b.kind, b.number, b.text) for b in docs[0].blocks] == [
        ("title", None, "Vyhláška města Horní Áíčkov o místních poplatcích"),
        ("article", "§ 1", "Předmět úpravy"),
        ("para", "1", "Město Horní Áíčkov vybírá místní poplatek ze psů."),
        ("para", "2", "Správcem poplatku je městský úřad."),
        ("article", "§ 2", ""),
        ("text", None, "Poplatek platí držitel psa."),
        ("article", "§ 3", "Účinnost"),
        ("text", None, "Tato vyhláška nabývá účinnosti dnem 1. ledna 2026."),
    ]
    assert [chunk.label for chunk in make_chunks(docs[0])] == [
        "§ 1 Předmět úpravy, odst. 1", "§ 1 Předmět úpravy, odst. 2", "§ 2", "§ 3 Účinnost"]


def test_clanek_bez_nazvu_si_nevezme_vetu_jako_nazev():
    """Za samotným „Čl. 6“ stojí rovnou text článku: věta se k nadpisu nepřilepí."""
    lines = [L("Čl. 6"), L("Toto nařízení nabývá účinnosti dnem vyhlášení."), L("Čl. 7"), L("Závěrečné ustanovení")]
    assert [line.text for line in join_wrapped(lines)] == [
        "Čl. 6", "Toto nařízení nabývá účinnosti dnem vyhlášení.", "Čl. 7 Závěrečné ustanovení"]


def test_cislovany_odstavec_a_datum():
    assert markers.paragraph("(2) Správcem poplatku je městský úřad.") == ("(1)", "2", "Správcem poplatku je městský úřad.")
    assert markers.paragraph("3. Poplatek je splatný do 31. března.") == ("1.", "3", "Poplatek je splatný do 31. března.")
    assert markers.paragraph("1. 6. 2026 nabývá vyhláška účinnosti.") is None       # datum, ne odstavec
    assert markers.letter("b) zahrady ve výši 2,0,") == ("b)", "zahrady ve výši 2,0,")


def test_priloha_a_cast():
    assert markers.appendix("Příloha č. 1")
    assert markers.appendix("Příloha č. 2 k obecně závazné vyhlášce č. 3/2025")
    assert markers.appendix("Příloha č. 1 této vyhlášky stanoví sazby poplatku.") is None
    assert markers.part("ČÁST PRVNÍ")
    assert markers.part("Část je splatná do konce roku.") is None


@pytest.mark.parametrize("text", ["Bc. Jana Nováková v. r.", "starosta", "starosta města místostarostka města",
                                  "…………………….     …………………", "Mgr. Petr Dvořák      Ing. Eva Malá", "1. místostarosta"])
def test_podpisove_radky(text):
    assert markers.is_signature(text)


@pytest.mark.parametrize("text", ["Správcem poplatku je městský úřad.", "Starosta může udělit výjimku.",
                                  "Žádost podává Ing. Novák do 15 dnů."])
def test_veta_neni_podpis(text):
    assert not markers.is_signature(text)


# ------------------------------------------------------------------ skládání zalomených řádků
def test_join_wrapped_slozi_odstavec_a_oddeli_znacky():
    lines = [L("(1) Poplatek se platí za psa staršího"), L("tří měsíců."), L("(2) Sazba činí 200 Kč."),
             Line(blank=True), L("Čl. 2", center=True), L("Splatnost", center=True)]
    joined = [line.text for line in join_wrapped(lines)]
    assert joined == ["(1) Poplatek se platí za psa staršího tří měsíců.", "(2) Sazba činí 200 Kč.", "Čl. 2 Splatnost"]


def test_join_wrapped_ulice_1_kvetna_neni_bod():
    """Řádek zalomený před „1. května“ nebo „17. listopadu“ nesmí začít nový bod."""
    lines = [L("(3) Obvod tvoří ulice Nádražní, Školní, Zahradní, Lipová,"), L("1. května, Gen. Svobody a"),
             L("17. listopadu.")]
    assert [line.text for line in join_wrapped(lines)] == [
        "(3) Obvod tvoří ulice Nádražní, Školní, Zahradní, Lipová, 1. května, Gen. Svobody a 17. listopadu."]


def test_join_wrapped_cislovane_body_za_sebou():
    lines = [L("Žadatel doloží:"), L("1. výpis z evidence,"), L("2. plnou moc,"), L("3. doklad o zaplacení.")]
    assert len(join_wrapped(lines)) == 4


# ------------------------------------------------------------------ rozbor na bloky
def _vyhlaska() -> list[Line]:
    return [
        L("MĚSTO HORNÍ ÁÍČKOV", center=True), L("Zastupitelstvo města Horní Áíčkov", center=True),
        L("Obecně závazná vyhláška města Horní Áíčkov", center=True), L("o místním poplatku ze psů", center=True),
        L("Zastupitelstvo města Horní Áíčkov se na svém zasedání dne 10. 12. 2025 usneslo vydat na základě § 14 zákona "
          "č. 565/1990 Sb., o místních poplatcích, tuto obecně závaznou vyhlášku:"),
        L("Čl. 1", center=True), L("Úvodní ustanovení", center=True),
        L("(1) Město Horní Áíčkov touto vyhláškou zavádí místní poplatek ze psů."),
        L("(2) Správcem poplatku je městský úřad."),
        L("Čl. 2", center=True), L("Sazba poplatku,", center=True), L("úlevy a osvobození", center=True),
        L("(1) Sazba poplatku za kalendářní rok činí:"),
        L("a) za prvního psa 300 Kč,"), L("b) za druhého a každého dalšího psa 600 Kč."),
        L("(2) Od poplatku je osvobozen držitel psa z útulku."),
        L("Čl. 3", center=True), L("Účinnost", center=True),
        L("Tato vyhláška nabývá účinnosti dnem 1. ledna 2026."),
        L("Bc. Jana Nováková v. r.     Mgr. Petr Dvořák v. r."), L("starostka     místostarosta"),
    ]


def test_build_blocks_vyhlaska():
    blocks, flags = build_blocks(_vyhlaska())
    kinds = [(b.kind, b.number) for b in blocks]
    assert kinds == [("title", None), ("text", None), ("article", "1"), ("para", "1"), ("para", "2"),
                     ("article", "2"), ("para", "1"), ("para", "2"), ("article", "3"), ("text", None)]
    assert blocks[0].text.splitlines() == ["MĚSTO HORNÍ ÁÍČKOV", "Zastupitelstvo města Horní Áíčkov",
                                           "Obecně závazná vyhláška města Horní Áíčkov", "o místním poplatku ze psů"]
    assert blocks[2].text == "Úvodní ustanovení"
    assert blocks[5].text == "Sazba poplatku, úlevy a osvobození"      # název článku na dvou řádcích
    assert [(i.marker, i.text) for i in blocks[6].items] == [("a)", "za prvního psa 300 Kč,"),
                                                             ("b)", "za druhého a každého dalšího psa 600 Kč.")]
    assert [b.anchor for b in blocks] == [f"b{n}" for n in range(len(blocks))]
    assert flags == ["podpisy"]
    assert "Nováková" not in " ".join(b.text for b in blocks)


def test_build_blocks_nazev_clanku_bez_napovedy():
    """Prostý text bez vystředění: název článku je jeden krátký řádek bez tečky za nadpisem."""
    blocks, _ = build_blocks([L("Čl. 1"), L("Předmět úpravy"), L("Tato vyhláška upravuje noční klid."),
                              L("Čl. 2"), L("Dobou nočního klidu se rozumí doba od 22 do 6 hodin.")])
    assert [(b.kind, b.text) for b in blocks] == [
        ("article", "Předmět úpravy"), ("text", "Tato vyhláška upravuje noční klid."),
        ("article", ""), ("text", "Dobou nočního klidu se rozumí doba od 22 do 6 hodin.")]


def test_build_blocks_vystredeny_kratky_clanek_neni_cely_nadpis():
    """Krátký článek bývá celý vystředěný; věta s tečkou do názvu nepatří. Jméno u funkce je podpis."""
    blocks, flags = build_blocks([L("Čl. 6", center=True), L("Účinnost", center=True),
                                  L("Toto nařízení nabývá účinnosti patnáctým dnem po vyhlášení.", center=True),
                                  L("Jan Novák", center=True), L("starosta", center=True)])
    assert [(b.kind, b.text) for b in blocks] == [("article", "Účinnost"),
                                                  ("text", "Toto nařízení nabývá účinnosti patnáctým dnem po vyhlášení.")]
    assert flags == ["podpisy"]


def test_build_blocks_pismena_bez_cisla_odstavce_a_podbody():
    blocks, _ = build_blocks([L("Čl. 3 Stejnokroj"), L("Stejnokroj strážníka tvoří:"), L("a) čepice,"), L("b) bunda:"),
                              L("zimní černá"), L("c) kalhoty.")])
    para = blocks[1]
    assert (para.kind, para.number, para.text) == ("para", None, "Stejnokroj strážníka tvoří:")
    assert [i.marker for i in para.items] == ["a)", "b)", "c)"]
    assert [c.text for c in para.items[1].children] == ["zimní černá"]


def test_build_blocks_veta_rozdelena_enterem_se_spoji():
    blocks, _ = build_blocks([L("Čl. 5 Sazba"), L("(1) Ceny za pronájem prodejního místa podle velikosti"),
                              L("stánku jsou uvedeny v ceníku."), L("(2) Poplatek je splatný ihned.")])
    assert blocks[1].text == "Ceny za pronájem prodejního místa podle velikosti stánku jsou uvedeny v ceníku."
    assert blocks[2].number == "2"


def test_build_blocks_jiny_druh_cislovani_jsou_body():
    """Odstavce „(1)“, pod písmenem body „1.“ – body patří k písmenu, ne mezi odstavce."""
    blocks, _ = build_blocks([L("Čl. 1 Slevy"), L("(1) Sleva ze vstupného se poskytuje:"), L("a) držitelům rodinné karty:"),
                              L("1. na koupaliště 20 Kč,"), L("2. do muzea 35 Kč,"), L("(2) Slevy se nesčítají.")])
    assert [b.number for b in blocks if b.kind == "para"] == ["1", "2"]
    assert [c.marker for c in blocks[1].items[0].children] == ["1.", "2."]


def test_build_blocks_priloha_a_tabulka():
    blocks, flags = build_blocks([L("Čl. 1 Sazby"), L("Sazby jsou v příloze č. 1."), L("Příloha č. 1"), L("Sazebník"),
                                  Line(rows=[["Druh", "Sazba"], ["stánek", "50 Kč"]])])
    assert [b.kind for b in blocks] == ["article", "text", "priloha", "table"]
    assert blocks[2].text == "Příloha č. 1 Sazebník"
    assert blocks[3].rows[1] == ["stánek", "50 Kč"]


def test_build_blocks_osamocena_priloha_bez_textu_se_ohlasi():
    """Příloha, ze které zbyl jen nadpis (mapa, obrázek), zůstává nadpisem a sestavení na ni upozorní."""
    blocks, flags = build_blocks([L("Čl. 1 Zákaz"), L("Konzumace alkoholu se zakazuje na místech vymezených v příloze č. 1."),
                                  L("Příloha č. 1 k obecně závazné vyhlášce o zákazu konzumace alkoholu")])
    assert [b.kind for b in blocks] == ["article", "text", "priloha"]
    assert "priloha-bez-textu" in flags


def test_build_blocks_dlouhy_vystredeny_nazev_a_veta_rozdelena_pred_paragrafem():
    title = ("Obecně závazná vyhláška města Horní Áíčkov, kterou se vymezují veřejná prostranství se zákazem pití alkoholu "
             "a stanoví výjimky pro městské slavnosti, trhy a předzahrádky restaurací")
    blocks, _ = build_blocks([L("MĚSTO HORNÍ ÁÍČKOV", center=True), L(title, center=True),
                              L("Zastupitelstvo města se usneslo vydat na základě ustanovení § 10 písm. a) a ustanovení"),
                              L("§ 84 odst. 2 písm. h) zákona č. 128/2000 Sb., o obcích, tuto obecně závaznou vyhlášku:"),
                              L("Čl. 1 Předmět"), L("Předmětem vyhlášky je zákaz konzumace alkoholu.")])
    assert blocks[0].kind == "title" and blocks[0].text.splitlines()[1] == title
    assert blocks[1].kind == "text" and blocks[1].text.endswith("tuto obecně závaznou vyhlášku:")
    assert [b.kind for b in blocks] == ["title", "text", "article", "text"]


def test_build_blocks_vycet_priloh_neni_priloha():
    blocks, _ = build_blocks([L("Čl. 1 Účinnost"), L("Nařízení nabývá účinnosti dnem vyhlášení."), L("Seznam příloh:"),
                              L("Příloha č. 1 – Plánek tržiště"), L("Příloha č. 2 – Ceník")])
    assert "priloha" not in [b.kind for b in blocks]


def test_build_blocks_obecny_dokument_s_nadpisy():
    blocks, flags = build_blocks([L("Návštěvní řád knihovny", style="h1"), L("Otevírací doba", style="h2"),
                                  L("Knihovna je otevřena v pracovní dny od 8 do 18 hodin."), L("Výpůjčky", style="h3"),
                                  L("Výpůjční doba knih je jeden měsíc.")])
    assert [b.kind for b in blocks] == ["title", "part", "text", "article", "text"]
    assert blocks[0].text == "Návštěvní řád knihovny"
    assert flags == []


def test_build_blocks_podpisy_v_tabulce():
    blocks, flags = build_blocks([L("Čl. 8 Účinnost"), L("Vyhláška nabývá účinnosti dnem 1. 1. 2026."),
                                  Line(rows=[["Bc. Jana Nováková v. r.", "Mgr. Petr Dvořák v. r."], ["starostka", "místostarosta"]])])
    assert [b.kind for b in blocks] == ["article", "text"]
    assert flags == ["podpisy"]


# ------------------------------------------------------------------ Word
def _docx(path, build) -> None:
    import docx

    document = docx.Document()
    build(document)
    document.save(str(path))


def test_docx_odstavce_tabulka_a_vystredeni(tmp_path):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    def build(document):
        title = document.add_paragraph("Obecně závazná vyhláška města Horní Áíčkov\no nočním klidu")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading = document.add_paragraph("Čl. 1")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name = document.add_paragraph("Doba nočního klidu")
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("Dobou nočního klidu se rozumí doba od dvacáté druhé do šesté hodiny.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Akce", "Termín"
        table.cell(1, 0).text, table.cell(1, 1).text = "Pouť", "srpen"

    path = tmp_path / "klid.docx"
    _docx(path, build)
    blocks, _ = build_blocks(read_document(path).lines)
    assert [(b.kind, b.number) for b in blocks] == [("title", None), ("article", "1"), ("text", None), ("table", None)]
    assert blocks[0].text == "Obecně závazná vyhláška města Horní Áíčkov\no nočním klidu"
    assert blocks[1].text == "Doba nočního klidu"
    assert blocks[3].rows == [["Akce", "Termín"], ["Pouť", "srpen"]]


def test_docx_automaticke_cislovani_se_dopocita(tmp_path):
    """Word čísla odstavců do textu nezapisuje – čtení je musí dopočítat z definice seznamu."""
    def build(document):
        document.add_paragraph("Čl. 1 Poplatek")
        document.add_paragraph("Město zavádí místní poplatek ze psů.", style="List Number")
        document.add_paragraph("Správcem poplatku je městský úřad.", style="List Number")

    path = tmp_path / "poplatek.docx"
    _docx(path, build)
    texts = [line.text for line in read_document(path).lines]
    assert texts[1].startswith("1") and texts[1].endswith("Město zavádí místní poplatek ze psů.")
    assert texts[2].startswith("2")
    blocks, _ = build_blocks(read_document(path).lines)
    assert [b.number for b in blocks if b.kind == "para"] == ["1", "2"]


def test_docx_pismena_s_vlastnim_seznamem_nevynuluji_cislo_odstavce(tmp_path):
    """Word pro písmena pod odstavcem založí nový seznam „číslovat od začátku“ a nový start zapíše pro všechny
    úrovně. Odstavec za písmeny musí pokračovat v řadě (1, 2, 3) – ne začít znovu od 1."""
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn

    def numbered(paragraph, num_id: int, level: int) -> None:
        num_pr = OxmlElement("w:numPr")
        for tag, value in (("w:ilvl", level), ("w:numId", num_id)):
            child = OxmlElement(tag)
            child.set(qn("w:val"), str(value))
            num_pr.append(child)
        paragraph._p.get_or_add_pPr().append(num_pr)

    def build(document):
        numbering = document.part.numbering_part.element
        numbering.append(parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90">'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="(%1)"/></w:lvl>'
            '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/></w:lvl>'
            '</w:abstractNum>'))
        numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="91"><w:abstractNumId w:val="90"/></w:num>'))
        numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="92"><w:abstractNumId w:val="90"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>'
            '<w:lvlOverride w:ilvl="1"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
        document.add_paragraph("Čl. 2 Koeficienty")
        for text, num_id, level in (("Koeficient pro stavby se stanoví takto:", 91, 0), ("Koeficient pro pozemky se stanoví takto:", 91, 0),
                                    ("zahrady 2,0,", 92, 1), ("sady 1,3.", 92, 1), ("Koeficienty platí od příštího roku.", 91, 0),
                                    ("Druhý seznam začíná znovu od jedné.", 92, 0), ("A pokračuje dvojkou.", 92, 0)):
            numbered(document.add_paragraph(text), num_id, level)

    path = tmp_path / "koeficienty.docx"
    _docx(path, build)
    marks = [line.text.split(" ", 1)[0] for line in read_document(path).lines[1:]]
    assert marks == ["(1)", "(2)", "a)", "b)", "(3)", "(1)", "(2)"]


def test_docx_obrazek_se_necte_ale_ohlasi(tmp_path):
    """Mapa vložená do Wordu jako obrázek: text se přečte, obrázek ne – a je to vidět ve výpisu i u znění."""
    import base64

    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")
    (tmp_path / "mapa.png").write_bytes(png)

    def build(document):
        document.add_paragraph("Čl. 1 Vymezení")
        document.add_paragraph("Zákaz platí na veřejných prostranstvích vyznačených v mapě.")
        document.add_picture(str(tmp_path / "mapa.png"))

    path = tmp_path / "alkohol.docx"
    _docx(path, build)
    assert read_document(path).flags == ["obrazky"]
    documents, _ = load_documents(tmp_path)
    assert "obrazky" in documents[0].flags


def test_docx_bez_textu_je_chyba(tmp_path):
    path = tmp_path / "prazdny.docx"
    _docx(path, lambda document: document.add_paragraph(""))
    with pytest.raises(NoTextError):
        read_document(path)


# ------------------------------------------------------------------ PDF
# znaky, které základní kódování PDF nemá, se do písma přidávají názvem (stejně to dělají kancelářské programy)
_PDF_EXTRA = {"Č": (200, "/Ccaron"), "č": (232, "/ccaron"), "ě": (236, "/ecaron"), "ř": (248, "/rcaron"),
              "ů": (249, "/uring"), "ň": (242, "/ncaron"), "Ř": (216, "/Rcaron")}


def _pdf(path, pages: list[list[tuple[int, int, str]]]) -> None:
    """Vyrobí PDF: každá strana je seznam (x, y, text). Písmo Helvetica s českými znaky přes tabulku rozdílů."""
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, DecodedStreamObject, DictionaryObject, NameObject, NumberObject

    def encode(text: str) -> bytes:
        raw = bytes(_PDF_EXTRA[ch][0] if ch in _PDF_EXTRA else ch.encode("cp1252")[0] for ch in text)
        return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")

    differences = ArrayObject()
    for code, glyph in sorted(_PDF_EXTRA.values()):
        differences += [NumberObject(code), NameObject(glyph)]
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): DictionaryObject({NameObject("/Type"): NameObject("/Encoding"),
                                                  NameObject("/BaseEncoding"): NameObject("/WinAnsiEncoding"),
                                                  NameObject("/Differences"): differences})})
    writer = PdfWriter()
    for lines in pages:
        page = writer.add_blank_page(width=595, height=842)
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"".join(b"BT /F1 11 Tf %d %d Td (%s) Tj ET\n" % (x, y, encode(text)) for x, y, text in lines))
        page[NameObject("/Contents")] = writer._add_object(stream)
    with open(path, "wb") as handle:
        writer.write(handle)


def test_pdf_cesky_text_clanky_a_cisla_stran(tmp_path):
    path = tmp_path / "klid.pdf"
    _pdf(path, [
        [(170, 780, "Obecně závazná vyhláška města Horní Áíčkov o nočním klidu"),
         (275, 740, "Článek 1"), (245, 724, "Doba nočního klidu"),
         (72, 700, "(1) Dobou nočního klidu se rozumí doba od dvacáté druhé do šesté hodiny, pokud"),
         (72, 686, "tato vyhláška nestanoví jinak."),
         (72, 668, "(2) Výjimky stanoví článek 2 této vyhlášky."),
         (290, 40, "1")],
        [(275, 780, "Článek 2"), (265, 764, "Výjimky"),
         (72, 740, "Doba nočního klidu se nedodržuje v noci z 31. prosince na 1. ledna."),
         (290, 40, "2")],
    ])
    blocks, _ = build_blocks(read_document(path).lines)
    assert [(b.kind, b.number, b.page) for b in blocks] == [
        ("title", None, 1), ("article", "1", 1), ("para", "1", 1), ("para", "2", 1), ("article", "2", 2), ("text", None, 2)]
    assert blocks[1].text == "Doba nočního klidu"
    # řádek zalomený uprostřed věty je zpět v jednom odstavci a číslo strany v textu není
    assert blocks[2].text == ("Dobou nočního klidu se rozumí doba od dvacáté druhé do šesté hodiny, pokud tato vyhláška "
                              "nestanoví jinak.")
    assert blocks[5].text.endswith("na 1. ledna.")


def test_pdf_bez_textu_je_chyba(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with (tmp_path / "sken.pdf").open("wb") as handle:
        writer.write(handle)
    with pytest.raises(NoTextError):
        read_document(tmp_path / "sken.pdf")


# ------------------------------------------------------------------ Markdown
MARKDOWN = """---
nazev: o místním poplatku ze psů
druh: Obecně závazná vyhláška
cislo: 5/2025
---

# Obecně závazná vyhláška města Horní Áíčkov o místním poplatku ze psů

Zastupitelstvo města Horní Áíčkov vydává tuto
obecně závaznou vyhlášku:

### Čl. 1 Sazba

1. Sazba poplatku činí:
    1. za prvního psa **300 Kč**,
    2. za dalšího psa 600 Kč.
2. Poplatek je splatný do 31. března.

| Druh | Sazba |
|---|---|
| pes v bytě | 300 Kč |
"""


def test_markdown_zahlavi_nadpisy_a_vnorene_seznamy(tmp_path):
    path = tmp_path / "psi.md"
    path.write_text(MARKDOWN, encoding="utf-8")
    extracted = read_document(path)
    assert extracted.meta == {"nazev": "o místním poplatku ze psů", "druh": "Obecně závazná vyhláška", "cislo": "5/2025"}
    blocks, _ = build_blocks(extracted.lines)
    assert [(b.kind, b.number) for b in blocks] == [("title", None), ("text", None), ("article", "1"), ("para", "1"),
                                                    ("para", "2"), ("table", None)]
    assert blocks[1].text == "Zastupitelstvo města Horní Áíčkov vydává tuto obecně závaznou vyhlášku:"
    assert [i.text for i in blocks[3].items] == ["za prvního psa 300 Kč,", "za dalšího psa 600 Kč."]
    assert blocks[5].rows == [["Druh", "Sazba"], ["pes v bytě", "300 Kč"]]


def test_vzor_prepisu_z_navodu_se_precte_tak_jak_navod_slibuje(tmp_path):
    """Vzor přepisu v dokumenty/README.md je to, podle čeho lidé i agenti píšou – musí se číst správně."""
    from conftest import ROOT

    readme = (ROOT / "dokumenty" / "README.md").read_text(encoding="utf-8")
    sample = readme.split("```markdown\n", 1)[1].split("\n```", 1)[0]
    path = tmp_path / "ozv-2007-03-o-zrizeni-mestske-policie.md"
    path.write_text(sample, encoding="utf-8")
    documents, skipped = load_documents(tmp_path)
    assert not skipped and len(documents) == 1
    doc = documents[0]
    assert doc.meta.label == "Obecně závazná vyhláška č. 3/2007 o zřízení městské policie"
    assert doc.meta.ucinnost_od == "2007-08-01"
    assert [(b.kind, b.number) for b in doc.blocks] == [
        ("title", None), ("text", None), ("article", "1"), ("para", "1"), ("para", "2"), ("article", "2"),
        ("para", None), ("priloha", None), ("table", None)]
    assert [i.marker for i in doc.blocks[6].items] == ["a)", "b)"]
    assert doc.blocks[7].text == "Příloha č. 1 Vyobrazení stejnokroje"


def test_markdown_udaje_tucnymi_radky(tmp_path):
    path = tmp_path / "smernice.md"
    path.write_text("**Druh:** Směrnice  \n**Číslo:** 3/2026  \n**Účinnost od:** 1. 6. 2026\n\n# Pravidla pro půjčování kol\n\n"
                    "### Čl. 1 Předmět\n\n1. Směrnice upravuje půjčování obecních kol.\n", encoding="utf-8")
    extracted = read_document(path)
    assert extracted.meta["druh"] == "Směrnice" and extracted.meta["číslo"] == "3/2026"


# ------------------------------------------------------------------ složka
def _write(path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def test_load_documents_prehled_prilohy_a_preskocene(tmp_path):
    _write(tmp_path / "sbirka" / "ozv-2026-02-o-nocnim-klidu.txt",
           "Obecně závazná vyhláška města Horní Áíčkov o nočním klidu\n\nČl. 1\nDoba nočního klidu\n\n"
           "Dobou nočního klidu se rozumí doba od dvacáté druhé do šesté hodiny.\n")
    _write(tmp_path / "sbirka" / "ozv-2026-02-o-nocnim-klidu-priloha-1.txt",
           "Seznam akcí\n\nPouť se koná v srpnu, noční klid je tu noc zkrácen.\n")
    _write(tmp_path / "sbirka" / "_neplatne" / "stara.txt", "Tohle už neplatí a číst se to nesmí, ani kdyby to bylo dlouhé.")
    _write(tmp_path / "README.md", "Sem patří dokumenty. Tento soubor se nečte, ani když má dost textu na dokument.")
    _write(tmp_path / "tabulka.xlsx", "x")
    _write(tmp_path / "prehled.json", json.dumps({"dokumenty": {"sbirka/ozv-2026-02-o-nocnim-klidu.txt": {
        "nazev": "o nočním klidu", "druh": "Obecně závazná vyhláška", "cislo": "2/2026", "ucinnost_od": "2026-04-29",
        "zdroj": "https://sbirkapp.gov.cz/detail/SPPAAA", "poznamka": "Ruší předpis č. 2/2025.",
        "prilohy": ["sbirka/ozv-2026-02-o-nocnim-klidu-priloha-1.txt"]}}}, ensure_ascii=False))
    documents, skipped = load_documents(tmp_path)
    assert len(documents) == 1
    doc = documents[0]
    assert doc.meta.label == "Obecně závazná vyhláška č. 2/2026 o nočním klidu"
    assert doc.meta.zdroj == "https://sbirkapp.gov.cz/detail/SPPAAA" and doc.meta.ucinnost_od == "2026-04-29"
    assert doc.meta.id == doc_id("sbirka/ozv-2026-02-o-nocnim-klidu")
    kinds = [b.kind for b in doc.blocks]
    assert kinds == ["title", "article", "text", "priloha", "text", "text"]
    assert doc.blocks[3].text == "Příloha č. 1"
    assert [b.anchor for b in doc.blocks] == ["b0", "b1", "b2", "b3", "b4", "b5"]
    assert [(s.path, s.reason[:6]) for s in skipped] == [("tabulka.xlsx", "formát")]


def test_load_documents_prepis_v_markdownu_ma_prednost(tmp_path):
    (tmp_path / "rad.pdf").write_bytes(b"%PDF-1.4 tohle neni citelne pdf")
    _write(tmp_path / "rad.md", "# Tržní řád\n\n### Čl. 1 Tržiště\n\n1. Tržiště je na náměstí a je otevřeno v sobotu.\n")
    documents, skipped = load_documents(tmp_path)
    assert [d.meta.soubor for d in documents] == ["rad.md"]
    assert documents[0].flags.count("prepis") == 1 and not skipped
    assert documents[0].meta.id == doc_id("rad")
    assert documents[0].meta.nazev == "Tržní řád"


def test_load_documents_necitelny_soubor_neshodi_ostatni(tmp_path):
    (tmp_path / "sken.pdf").write_bytes(b"%PDF-1.4 tohle neni citelne pdf")
    _write(tmp_path / "dobry.txt", "Provozní řád hřiště\n\nHřiště je otevřeno denně od 8 do 20 hodin, vstup je volný.\n")
    documents, skipped = load_documents(tmp_path)
    assert [d.meta.soubor for d in documents] == ["dobry.txt"]
    assert [s.path for s in skipped] == ["sken.pdf"]


def test_load_documents_prazdna_nebo_chybejici_slozka(tmp_path):
    assert load_documents(tmp_path / "neni") == ([], [])


# ------------------------------------------------------------------ nálezy ze zkoušky na předpisech jiné obce
@pytest.mark.parametrize("text", [
    "starosta města Horního Áíčkova", "místostarosta města Horního Áíčkova starosta města Horního Áíčkova", "starostka obce Horní Lhota",
    "primátor statutárního města Ústí nad Labem", "zástupce starosty města",
    "Jan Novák místostarosta města Horního Áíčkova", "Jana Nováková místostarostka",
    "Vyvěšeno na úřední desce dne: 12. 3. 2021", "Sejmuto z úřední desky dne:", "Sňato dne 28. 3. 2021",
])
def test_podpis_s_nazvem_obce_jmenem_nebo_dolozkou_o_vyveseni(text):
    assert markers.is_signature(text)


@pytest.mark.parametrize("text", [
    "Starosta města Horního Áíčkova může udělit výjimku", "Starosta města rozhoduje o výjimkách do 30 dnů",
    "Jan Novák podal žádost starostovi", "Vyvěšení na úřední desce trvá patnáct dnů.",
])
def test_veta_o_starostovi_neni_podpis(text):
    assert not markers.is_signature(text)


@pytest.mark.parametrize("text", [
    "Příloha č. 1, vyhlášky č. 4/2022 o místním poplatku ze psů",
    "Příloha č. 2 vyhlášky č. 4/2022 o místním poplatku ze psů",
    "Příloha č. 1 obecně závazné vyhlášky č. 2/2024",
    "Příloha č. 3 nařízení č. 1/2025",
])
def test_priloha_s_odkazem_na_predpis_je_nadpis(text):
    assert markers.appendix(text)


def test_datum_neni_cislo_odstavce():
    assert markers.paragraph("28. října. Medaile pak vítězům předá pořadatel závodu.") is None
    assert markers.paragraph("1. května se koná průvod.") is None
    assert markers.paragraph("3. Poplatek je splatný do 31. března.") == ("1.", "3", "Poplatek je splatný do 31. března.")


def test_build_blocks_datum_na_zacatku_radku_je_pokracovani_vety():
    """Word: autor uprostřed věty stiskl Enter před datem – „28. října“ není odstavec 28."""
    blocks, _ = build_blocks([L("Čl. 7 Běh městem"), L("(1) Závod se běhá každý rok ve sváteční den"),
                              L("28. října. Medaile předá vítězům pořadatel.")])
    assert [(b.kind, b.number) for b in blocks] == [("article", "7"), ("para", "1")]
    assert blocks[1].text == "Závod se běhá každý rok ve sváteční den 28. října. Medaile předá vítězům pořadatel."


def test_build_blocks_radek_za_carkou_pokracuje_i_velkym_pismenem():
    """Výčet ulic zalomený do dalšího odstavce: řádek začíná velkým písmenem, ale věta před ním neskončila."""
    blocks, _ = build_blocks([L("Čl. 4 Parkovné"), L("(1) Parkovné za den činí:"),
                              L("a) 60 Kč v ulicích Nádražní, Školní,"), L("Zahradní a Lipová,"),
                              L("b) 25 Kč na parkovišti u nádraží.")])
    para = blocks[1]
    assert [item.marker for item in para.items] == ["a)", "b)"]
    assert para.items[0].text == "60 Kč v ulicích Nádražní, Školní, Zahradní a Lipová,"


def test_clanek_s_chybou_cteni_skenu_a_malym_pismenem():
    assert markers.article("ČI. 10") == ("10", "")            # čtení skenu: velké I místo malého l
    assert markers.article("článek 3") == ("3", "")
    assert markers.article("článek 3 této vyhlášky se nepoužije") is None


def test_pdf_nadpis_clanku_nahore_na_strance_neni_zahlavi():
    """Když tři strany z pěti začínají „Čl. N“, nesmí se to vzít jako opakované záhlaví a zahodit.
    Záhlaví, zápatí a čísla stran se hledají jen na okrajích stránky – číslo uprostřed textu zůstává."""
    from chatbot.ingest.extract import lines_from_pages

    pages = [
        ["Vyhláška města Horní Áíčkov", "", "Čl. 1", "Úvod", "", "Vyhláška upravuje místní poplatek ze psů.", "", "Sbírka předpisů města", "1"],
        ["Čl. 3", "Sazba", "", "(1) Sazba činí ročně", "150", "korun za jednoho psa.", "", "Sbírka předpisů města", "2"],
        ["Čl. 5", "Osvobození", "", "Od poplatku je osvobozen držitel průkazu.", "", "Sbírka předpisů města", "3"],
        ["Čl. 10", "Účinnost", "", "Vyhláška nabývá účinnosti dnem vyhlášení.", "", "Sbírka předpisů města", "4"],
    ]
    texts = [line.text for line in lines_from_pages(pages) if not line.blank]
    assert [text for text in texts if text.startswith("Čl.")] == ["Čl. 1 Úvod", "Čl. 3 Sazba", "Čl. 5 Osvobození", "Čl. 10 Účinnost"]
    assert "(1) Sazba činí ročně 150 korun za jednoho psa." in texts
    assert not any(text in ("1", "2", "3", "4") or text.startswith("Sbírka předpisů") for text in texts)


def test_prostrkany_nadpis_o_jednom_slove():
    from chatbot.ingest.extract import _unspace

    assert _unspace("N a ř í z e n í") == "Nařízení"
    assert _unspace("M ě s t o   H o r n í   Á í č k o v") == "Město Horní Áíčkov"
    assert _unspace("a b c") == "a b c"                         # tři písmena nejsou prostrkaný nadpis


def test_dlouha_priloha_v_jednom_bloku_zustane_textem(tmp_path):
    """Příloha, jejíž celý obsah je jeden dlouhý odstavec začínající slovem „Příloha“, se nesmí celá stát
    nadpisem – nedala by se pak najít."""
    from chatbot.index.chunks import make_chunks

    _write(tmp_path / "rad.txt", "Požární řád města Horní Áíčkov\n\nČl. 1\nÚvod\n\nŘád upravuje požární ochranu ve městě.\n")
    _write(tmp_path / "rad-priloha-1.txt",
           "Příloha č. 4 Zdroje vody pro hašení: " + "hydrant v ulici Nádražní, požární nádrž u školy, " * 12 + "rybník.\n")
    documents, skipped = load_documents(tmp_path)
    assert not skipped
    blocks = documents[0].blocks
    assert [block.kind for block in blocks][-2:] == ["priloha", "text"]
    assert blocks[-2].text == "Příloha č. 1" and "hydrant" in blocks[-1].text
    assert any("hydrant" in chunk.text for chunk in make_chunks(documents[0]))


def test_kratky_udaj_s_cislem_je_usek_samotne_cislo_ne():
    from chatbot.index.chunks import make_chunks
    from chatbot.ingest.model import Block, DocMeta, ParsedDoc

    doc = ParsedDoc(meta=DocMeta(id=1, soubor="rad.txt", nazev="Požární řád"), blocks=[
        Block(kind="priloha", text="Příloha č. 2", anchor="b0"), Block(kind="text", text="150 nebo 112", anchor="b1"),
        Block(kind="text", text="6", anchor="b2"), Block(kind="text", text="viz výše", anchor="b3")])
    # samotné číslo se zahodí; krátký dovětek se slovy se připojí k úseku před sebou (sám úsekem není)
    assert [chunk.text for chunk in make_chunks(doc)] == ["150 nebo 112\nviz výše"]


def test_podpisy_uprostred_dokumentu_pred_cenikem():
    """Za podpisy ještě následuje ceník bez slova „Příloha“ – podpisy přesto do textu nepatří."""
    lines = [L("Čl. 10 Závěrečná ustanovení"), L("Toto nařízení nabývá účinnosti dnem 1. ledna 2026."),
             L("Mgr. Jan Novák v. r."), L("starosta města Horního Áíčkova"), L("Petr Dvořák v. r."), L("místostarosta města Horního Áíčkova"),
             L("Ceník"), L("k nařízení o půjčovně lodiček"), L("Půjčení lodičky stojí 80 Kč za hodinu, na celý den 450 Kč.")]
    blocks, flags = build_blocks(lines)
    text = " ".join(block.text for block in blocks)
    assert "Novák" not in text and "Dvořák" not in text and "starosta" not in text
    assert "Ceník" in text and "450 Kč" in text and flags == ["podpisy"]


def test_vycet_funkci_uprostred_textu_neni_podpis():
    lines = [L("Čl. 2 Komise"), L("Členy komise jsou:"), L("starosta"), L("místostarosta"), L("tajemník"),
             L("Čl. 3 Účinnost"), L("Nařízení nabývá účinnosti dnem vyhlášení.")]
    blocks, _ = build_blocks(lines)
    assert "místostarosta" in " ".join(block.text for block in blocks)


def test_poznamky_pod_carou_uprostred_vyctu_vycet_neprerusi():
    """PDF: výčet pokračuje na další straně a mezi písmeny b) a c) stojí poznámky pod čarou z konce stránky."""
    blocks, _ = build_blocks([L("Čl. 2 Pořadí důležitosti"), L("(2) Chodníky se rozdělují na:"),
                              L("a) I. pořadí – chodníky v centru,"), L("b) II. pořadí – spojovací chodníky,"),
                              L("1/ dopravní značka s dodatkovou tabulkou"), L("2/ nařízení města o neudržovaných úsecích"),
                              L("c) neudržované."), L("(3) Úseky vymezuje příloha.")])
    para = blocks[1]
    assert (para.kind, para.number) == ("para", "2") and [item.marker for item in para.items] == ["a)", "b)", "c)"]
    assert para.items[1].text == "II. pořadí – spojovací chodníky,"
    assert [block.text for block in blocks if block.kind == "text"] == ["1/ dopravní značka s dodatkovou tabulkou",
                                                                         "2/ nařízení města o neudržovaných úsecích"]
    assert blocks[-1].number == "3"


def test_poznamka_pod_carou_hned_za_navetim_neprevezme_pismena():
    """Poznámka pod čarou stojí hned za návětím s dvojtečkou: písmena za ní patří dál k odstavci, ne k poznámce."""
    blocks, _ = build_blocks([L("Čl. 2 Pořadí"), L("(2) Chodníky se rozdělují na:"), L("1/ zákon o pozemních komunikacích"),
                              L("a) hlavní,"), L("b) vedlejší."), L("(3) Úseky vymezuje příloha.")])
    para = blocks[1]
    assert (para.kind, para.number, [item.marker for item in para.items]) == ("para", "2", ["a)", "b)"])
    assert [block.text for block in blocks if block.kind == "text"] == ["1/ zákon o pozemních komunikacích"]


# ------------------------------------------------------------------ PDF: rozbité kódování, mapy, okraje stran
def test_pdf_text_po_dvojicich_bajtu_se_prevede_zpet():
    """Mapy vyvedené z GIS vracejí popisky po dvojicích bajtů UTF-16 – do textu se tak dostávaly znaky NUL."""
    from chatbot.ingest.extract import clean_page_text, decode_byte_pairs

    def pairs(text: str) -> str:
        return text.encode("utf-16-be").decode("latin-1")

    broken = pairs("Příloha č. 1 – mapa zón") + "\n   " + pairs("Zóna 2") + "   0 125 250 metry"
    assert clean_page_text(broken).split("\n") == ["Příloha č. 1 – mapa zón", "   Zóna 2   0 125 250 metry"]
    assert decode_byte_pairs("Běžný text bez vady.") == "Běžný text bez vady."
    assert "\x00" not in clean_page_text("a\x00b\x00c\x00d\x00e")       # co po převodu zbude, do textu nesmí
    assert clean_page_text("konec\x0cstrany") == "konec strany"


def test_pdf_strana_jen_s_popisky_je_mapa():
    from chatbot.ingest.extract import labels_only

    parcels = [f"{1000 + i}/{i % 7 + 1}   {2000 + i}" for i in range(60)]
    assert labels_only(["Příloha č. 1 k nařízení o placeném stání"] + parcels + ["Placené zóny", "0 125 250 metry"])
    # sazebník má u čísel slova – mapa to není
    assert not labels_only([f"{i}. za stánek do {i} m2 denně {10 * i} Kč" for i in range(1, 40)])
    assert not labels_only(["12/3 45/6"])                       # pár čísel na skoro prázdné straně se neposuzuje


def _map_labels() -> list[tuple[int, int, str]]:
    return [(50 + (i % 10) * 50, 760 - (i // 10) * 12, f"{1000 + i}/{i % 7 + 1}") for i in range(120)]


def test_pdf_priloha_s_mapou_necha_jen_nadpis_a_ohlasi_se(tmp_path):
    """Z mapy jdou vytáhnout jen čísla parcel. Jako text přílohy by mátla – zůstane nadpis a upozornění."""
    from chatbot.index.chunks import make_chunks

    heading = "Příloha č. 1 k nařízení města Horní Áíčkov o placeném stání"
    _write(tmp_path / "stani.txt", "Nařízení města Horní Áíčkov o placeném stání\n\nČl. 1\nZóny\n\n"
                                   "Placené stání je v zónách vyznačených v příloze č. 1.\n")
    _pdf(tmp_path / "stani-priloha-1.pdf", [[(72, 800, heading)] + _map_labels()])
    extracted = read_document(tmp_path / "stani-priloha-1.pdf")
    assert [line.text for line in extracted.lines] == [heading] and extracted.flags == ["popisky"]
    documents, skipped = load_documents(tmp_path)
    assert not skipped
    doc = documents[0]
    assert (doc.blocks[-1].kind, doc.blocks[-1].text) == ("priloha", heading)
    assert "popisky" in doc.flags and "priloha-bez-textu" in doc.flags
    assert not any("1000/" in chunk.text for chunk in make_chunks(doc))


def test_pdf_jen_mapa_bez_nadpisu_se_necte_a_hlaska_rika_proc(tmp_path):
    from chatbot.ingest.extract import LabelsOnlyError

    _pdf(tmp_path / "mapa.pdf", [_map_labels()])
    with pytest.raises(LabelsOnlyError):
        read_document(tmp_path / "mapa.pdf")
    documents, skipped = load_documents(tmp_path)
    assert not documents and "mapa nebo výkres" in skipped[0].reason


def test_pdf_cislo_strany_s_pomlckami_se_zahodi_a_veta_pres_stranu_se_spoji():
    """Zápatí „- 2 -“ vypadá jako odrážka. Nesmí zůstat v textu ani rozdělit větu, která pokračuje na další straně."""
    from chatbot.ingest.extract import lines_from_pages

    pages = [["Čl. 1 Parkovné", "(1) Parkovné činí za hodinu dvacet korun a za celý den", "- 1 -"],
             ["sto korun.", "(2) Parkovné se platí v automatu.", "- 2 -"],
             ["(3) Držitelé průkazu ZTP parkují zdarma.", "- 3 -"]]
    assert [line.text for line in lines_from_pages(pages)] == [
        "Čl. 1 Parkovné", "(1) Parkovné činí za hodinu dvacet korun a za celý den sto korun.",
        "(2) Parkovné se platí v automatu.", "(3) Držitelé průkazu ZTP parkují zdarma."]


def test_pdf_cislo_na_zacatku_strany_ktere_neni_cislem_strany_zustava():
    from chatbot.ingest.extract import lines_from_pages

    pages = [["Čl. 1 Vstupné", "Vstupné pro dospělé činí", "1"], ["100", "korun, pro děti polovinu.", "2"]]
    assert [line.text for line in lines_from_pages(pages)] == [
        "Čl. 1 Vstupné", "Vstupné pro dospělé činí 100 korun, pro děti polovinu."]


def test_pdf_radky_lisici_se_jen_cisly_nejsou_zahlavi():
    """Seznam parcel, který vyšel na okraje stran: řádky se liší jen čísly, ale záhlaví to není."""
    from chatbot.ingest.extract import lines_from_pages

    pages = [[f"p. č. {100 * page + row}/2 – chodník" for row in range(1, 9)] for page in range(1, 5)]
    text = " ".join(line.text for line in lines_from_pages(pages))
    assert all(f"p. č. {100 * page + row}/2" in text for page in range(1, 5) for row in range(1, 9))


def test_pdf_opakovane_zahlavi_a_strana_n_z_m_se_zahodi():
    from chatbot.ingest.extract import lines_from_pages

    pages = [["Vyhláška č. 5/2025", f"Strana {page} z 3", f"({page}) Odstavec číslo {page} má svůj text."]
             for page in range(1, 4)]
    assert [line.text for line in lines_from_pages(pages)] == [
        f"({page}) Odstavec číslo {page} má svůj text." for page in range(1, 4)]


def test_prostrkana_cisla_se_neslepi():
    from chatbot.ingest.extract import _unspace

    assert _unspace("1 2   3 4") == "1 2 3 4"                   # řada čísel – slepením by vznikla jiná čísla
    assert _unspace("1 5 0   n e b o   1 1 2") == "150 nebo 112"  # prostrkaný text s čísly se složí


# ------------------------------------------------------------------ nadpis článku u levého okraje
def test_join_wrapped_text_za_nadpisem_s_nazvem_je_novy_odstavec():
    """Prostý text a PDF bez vystředěných nadpisů: článek se nesmí slít s nadpisem a nadpis tím zmizet."""
    lines = [L("Čl. 2"), L("Účinnost"), L("Tato vyhláška nabývá účinnosti patnáctým dnem po dni vyhlášení a platí"),
             L("do odvolání."), L("Čl. 3 Zrušovací ustanovení"), L("Zrušuje se vyhláška o místním poplatku z roku 2011"),
             L("včetně jejích příloh.")]
    assert [line.text for line in join_wrapped(lines)] == [
        "Čl. 2 Účinnost", "Tato vyhláška nabývá účinnosti patnáctým dnem po dni vyhlášení a platí do odvolání.",
        "Čl. 3 Zrušovací ustanovení", "Zrušuje se vyhláška o místním poplatku z roku 2011 včetně jejích příloh."]


def test_join_wrapped_dlouhy_radek_za_samotnou_znackou_neni_nazev():
    """Za samotným „Čl. 6“ začíná rovnou zalomený odstavec: jeho první řádek (bez tečky) není název článku."""
    lines = [L("Čl. 5"), L("Evidence"),
             L("(1) Provozovatel tržiště vede seznam prodejců a na vyžádání ho předloží"), L("kontrole městského úřadu."),
             L("Čl. 6"), L("Toto nařízení nabývá účinnosti patnáctým dnem po dni vyhlášení ve Sbírce"),
             L("právních předpisů.")]
    joined = [line.text for line in join_wrapped(lines)]
    assert joined[0] == "Čl. 5 Evidence" and joined[2] == "Čl. 6"
    assert joined[3] == "Toto nařízení nabývá účinnosti patnáctým dnem po dni vyhlášení ve Sbírce právních předpisů."


def test_join_wrapped_vystredena_veta_za_nazvem_clanku_neni_nazev():
    lines = [L("Čl. 6", center=True), L("Účinnost", center=True),
             L("Toto nařízení nabývá účinnosti dnem vyhlášení.", center=True)]
    assert [line.text for line in join_wrapped(lines)] == ["Čl. 6 Účinnost", "Toto nařízení nabývá účinnosti dnem vyhlášení."]


# ------------------------------------------------------------------ podpis × obsah
def test_udaj_se_jmenem_a_titulem_neni_podpis():
    """Na konci přílohy stojí kontakt. Jméno s titulem v něm je, ale podpis to není – patří k obsahu."""
    assert not markers.is_signature("Velitel jednotky: Ing. Petr Dvořák, tel. 555 123 456")
    assert not markers.is_signature("Žádosti přijímá Mgr. Jana Nováková v kanceláři č. 12")
    assert markers.is_signature("Mgr. Jana Nováková") and markers.is_signature("Ing. Petr Dvořák, starosta")
    assert markers.is_signature("Ing. Petr Dvořák, Ph.D., MBA")
    blocks, flags = build_blocks([L("Čl. 1 Požární ochrana"), L("Jednotku svolává velitel."),
                                  L("Příloha č. 1 Kontakty"), L("Velitel jednotky: Ing. Petr Dvořák, tel. 555 123 456")])
    assert "Velitel jednotky" in blocks[-1].text and "podpisy" not in flags


@pytest.mark.parametrize("text", [
    "……………………. ………………… Mgr. Jan Novák Ing. Petr Dvořák místostarosta starosta",     # sloupce podpisů slité z PDF
    "Jana Müllerová Mgr. Petr Dvořák", "Ing. Ľubomír Wójcik", "starosta Ing. Jan Novák",
    "Zpracovala mjr. Ing. Jana Nováková dne 12. února 2026", "Vypracoval: Jan Novák",
])
def test_podpis_s_teckami_cizimi_pismeny_nebo_zpracovatelem(text):
    assert markers.is_signature(text)


def test_jmeno_s_prehlaskou_u_funkce_je_podpis():
    assert markers.is_bare_name("Jana Müllerová")
    blocks, flags = build_blocks([L("Čl. 9 Účinnost"), L("Vyhláška nabývá účinnosti dnem vyhlášení."),
                                  L("Jana Müllerová"), L("místostarostka města Horního Áíčkova")])
    assert [b.kind for b in blocks] == ["article", "text"] and flags == ["podpisy"]


@pytest.mark.parametrize("text", ["Zpracoval odbor financí městského úřadu.", "Schválilo zastupitelstvo města dne 10. 12. 2025"])
def test_veta_o_zpracovani_bez_jmena_neni_podpis(text):
    assert not markers.is_signature(text)


def test_vycet_funkci_pred_prilohou_a_nazev_clanku_tajemnik_nejsou_podpis():
    lines = [L("Čl. 8 Komise"), L("Členy komise jsou:"), L("starosta"), L("místostarosta"), L("tajemník"),
             L("Příloha č. 1 Jednací řád"), L("Komise jedná podle potřeby, nejméně dvakrát do roka.")]
    text = " ".join(block.text for block in build_blocks(lines)[0])
    assert "starosta" in text and "tajemník" in text
    blocks, flags = build_blocks([L("Čl. 1 Starosta"), L("Starosta zastupuje město navenek."), L("Čl. 2"), L("Tajemník"),
                                  L("Tajemník řídí městský úřad.")])
    assert (blocks[2].kind, blocks[2].text) == ("article", "Tajemník") and "podpisy" not in flags


def test_dvouslovny_nazev_nad_podpisem_zustava():
    """„Dolní Lhota“ vypadá jako jméno a příjmení. Podpisem je samotné jméno jen hned u řádku s funkcí."""
    lines = [L("Čl. 3 Místní části"), L("Vyhláška platí v místních částech:"), L("Horní Lhota"), L("Dolní Lhota"),
             L("Mgr. Jana Áíčková v. r."), L("starostka")]
    blocks, flags = build_blocks(lines)
    text = " ".join(block.text for block in blocks)
    assert "Dolní Lhota" in text and "Áíčková" not in text and "starostka" not in text and flags == ["podpisy"]


@pytest.mark.parametrize("text, expected", [
    ("Čl. 4 Odstraňování závad ve schůdnosti", ("4", "Odstraňování závad ve schůdnosti")),
    ("Čl. 7 Odstavování vozidel", ("7", "Odstavování vozidel")),
    ("Čl. 3 Bodové hodnocení žádostí", ("3", "Bodové hodnocení žádostí")),
    ("Čl. 2 Nařízení odstranění stavby", ("2", "Nařízení odstranění stavby")),
    ("§ 5 Odstraňování odpadu", ("§ 5", "Odstraňování odpadu")),
])
def test_nazev_clanku_ktery_zacina_jako_odkaz_je_nadpis(text, expected):
    """„Odst…“, „Bod…“ a „Nařízení…“ na začátku názvu nejsou odkaz „odst. 2“, „bod 3“, „nařízení č. …“."""
    assert markers.article(text) == expected


@pytest.mark.parametrize("text", ["článek 5.", "Čl. 3 se mění takto:", "Čl. 5 se zrušuje", "Čl. 4 odstavce 2 a 3 se zrušují",
                                  "Čl. 6 bodu 3 se nepoužije", "Čl. 2 zní:"])
def test_veta_o_clanku_neni_nadpis(text):
    assert markers.article(text) is None


@pytest.mark.parametrize("text", ["Část Horní Lhota – celé území", "Část obce Horní Lhota",
                                  "Část poplatku lze uhradit ve splátkách", "Hlava rodiny", "Oddíl kultury městského úřadu"])
def test_veta_zacinajici_slovem_cast_neni_nadpis_casti(text):
    assert markers.part(text) is None


@pytest.mark.parametrize("text", ["ČÁST PRVNÍ", "Část druhá", "Hlava II Společná ustanovení", "ODDÍL 3", "Část A – Poplatky",
                                  "Kapitola čtvrtá: Přechodná ustanovení"])
def test_nadpis_casti_s_cislem_pismenem_nebo_radovou_cislovkou(text):
    assert markers.part(text) == text


def test_txt_clanky_u_leveho_okraje_zustanou_clanky(tmp_path):
    from chatbot.index.chunks import make_chunks

    _write(tmp_path / "rad.txt", "Tržní řád města Horní Áíčkov\n\nČl. 1\nTržiště\nTržiště je na náměstí a prodává se na něm "
                                 "v sobotu dopoledne.\nČl. 2 Účinnost\nTento řád nabývá účinnosti dnem vyhlášení.\n")
    documents, _ = load_documents(tmp_path)
    assert documents[0].flags == []
    assert [chunk.label for chunk in make_chunks(documents[0])] == ["čl. 1 Tržiště", "čl. 2 Účinnost"]
